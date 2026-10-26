import io
import json
import math
import os
import urllib.error
import urllib.request
from datetime import date, datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE44_URL = "https://finsim-pay-flow.base44.app/api"
APP_DIR = Path(__file__).resolve().parent
DEBTS_FILE = APP_DIR / "debts.json"
HOST = "0.0.0.0"
PORT = int(os.environ.get("PORT", "10000"))


def fetch_base44_entity(entity_name):
    """Fetch permitted records from Base44 using a private token."""
    token = os.environ.get("BASE44_TOKEN")
    if not token:
        raise RuntimeError("BASE44_TOKEN is not configured.")
    if entity_name not in ("BankAccount", "Transaction"):
        raise ValueError("Unsupported entity.")
    url = f"{BASE44_URL}/entities/{entity_name}"
    request = urllib.request.Request(
        url,
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
        method="GET",
    )
    with urllib.request.urlopen(request, timeout=20) as response:
        data = json.loads(response.read().decode("utf-8"))
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ("items", "data", "results"):
            if isinstance(data.get(key), list):
                return data[key]
    raise ValueError(f"Unexpected Base44 response for {entity_name}.")


def load_debts():
    if not DEBTS_FILE.exists():
        return []
    try:
        with open(DEBTS_FILE, "r", encoding="utf-8") as file:
            data = json.load(file)
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        print("Warning: could not read debts.json.")
        return []


def save_debts(debts):
    with open(DEBTS_FILE, "w", encoding="utf-8") as file:
        json.dump(debts, file, indent=2, ensure_ascii=False)


def calculate_warning_status(amount: float, days_left: int) -> dict:
    HIGH_AMOUNT_THRESHOLD = 1000.0
    MEDIUM_AMOUNT_THRESHOLD = 250.0
    if days_left <= 2 and amount >= MEDIUM_AMOUNT_THRESHOLD:
        return {"level": "Critical", "color": "red", "hex": "#EF4444"}
    elif amount >= HIGH_AMOUNT_THRESHOLD or days_left <= 5:
        return {"level": "High", "color": "orange", "hex": "#F97316"}
    elif days_left <= 15 or amount >= MEDIUM_AMOUNT_THRESHOLD:
        return {"level": "Medium", "color": "yellow", "hex": "#EAB308"}
    return {"level": "Low", "color": "green", "hex": "#22C55E"}


def prepare_debt(debt):
    due_date = date.fromisoformat(debt["due_date"])
    days_left = (due_date - date.today()).days
    amount = float(debt["amount"])
    return {**debt, "days_left": days_left, "overdue": days_left < 0,
            "warning": calculate_warning_status(amount, days_left)}


def get_debts_with_warnings():
    debts = [prepare_debt(debt) for debt in load_debts()]
    priority = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
    debts.sort(key=lambda item: (priority[item["warning"]["level"]], item["days_left"]))
    return debts


def get_transaction_date(transaction):
    for field in ("timestamp", "date", "transaction_date", "created_date"):
        value = transaction.get(field)
        if not value:
            continue
        try:
            return datetime.fromisoformat(str(value).replace("Z", "+00:00")).date()
        except ValueError:
            try:
                return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
    return None


def create_monthly_excel(transactions, selected_month):
    filtered = []
    for transaction in transactions:
        tx_date = get_transaction_date(transaction)
        if tx_date and tx_date.strftime("%Y-%m") == selected_month:
            filtered.append((tx_date, transaction))
    columns = sorted({key for _, transaction in filtered for key in transaction.keys()})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Transactions"
    sheet.append([f"Transaction Report: {selected_month}"])
    sheet.append([f"Transactions found: {len(filtered)}"])
    sheet.append([])
    if columns:
        sheet.append(columns)
        for _, transaction in filtered:
            row = []
            for column in columns:
                value = transaction.get(column)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                row.append(value)
            sheet.append(row)
    else:
        sheet.append(["message"])
        sheet.append(["No transactions found for this month."])
    sheet["A1"].font = Font(bold=True, size=16)
    sheet["A2"].font = Font(italic=True)
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="245DB5")
    sheet.freeze_panes = "A5"
    if columns:
        sheet.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{sheet.max_row}"
    for cells in sheet.columns:
        letter = get_column_letter(cells[0].column)
        longest = max(len(str(cell.value or "")) for cell in cells)
        sheet.column_dimensions[letter].width = min(longest + 2, 35)
    summary = workbook.create_sheet("Summary")
    summary.append(["Monthly Report Summary"])
    summary.append(["Month", selected_month])
    summary.append(["Transaction count", len(filtered)])
    summary["A1"].font = Font(bold=True, size=16)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 28
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), len(filtered)


class AppHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(APP_DIR), **kwargs)

    def send_json(self, status, data):
        body = json.dumps(data, default=str).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def read_json_body(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise ValueError("Invalid content length.")
        if length <= 0 or length > 100_000:
            raise ValueError("Request body is empty or too large.")
        raw = self.rfile.read(length)
        payload = json.loads(raw.decode("utf-8"))
        if not isinstance(payload, dict):
            raise ValueError("Expected a JSON object.")
        return payload

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/health":
            self.send_json(200, {"status": "running", "app": "Matrix Finance Tracker "})
            return
        if path == "/api/dashboard":
            try:
                accounts = fetch_base44_entity("BankAccount")
                transactions = fetch_base44_entity("Transaction")
                self.send_json(200, {"accounts": accounts, "transactions": transactions})
            except urllib.error.HTTPError as exc:
                print(f"Base44 HTTP error: {exc.code}")
                self.send_json(502, {"error": "Base44 rejected the request. Check token validity and permissions."})
            except Exception as exc:
                print(f"Dashboard request failed: {exc}")
                self.send_json(500, {"error": "Could not load dashboard data. Check server logs for details."})
            return
        if path == "/api/debts":
            try:
                self.send_json(200, {"debts": get_debts_with_warnings()})
            except Exception as exc:
                print(f"Debt loading failed: {exc}")
                self.send_json(500, {"error": "Could not load saved debts."})
            return
        if path == "/api/monthly-report":
            query = parse_qs(urlparse(self.path).query)
            selected_month = query.get("month", [""])[0]
            try:
                datetime.strptime(selected_month, "%Y-%m")
            except ValueError:
                self.send_json(400, {"error": "Select a month in YYYY-MM format."})
                return
            try:
                transactions = fetch_base44_entity("Transaction")
                excel_bytes, count = create_monthly_excel(transactions, selected_month)
                filename = f"transactions_{selected_month}.xlsx"
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                self.send_header("Content-Length", str(len(excel_bytes)))
                self.send_header("Cache-Control", "no-store")
                self.end_headers()
                self.wfile.write(excel_bytes)
                print(f"Created {filename} with {count} transaction(s).")
            except urllib.error.HTTPError as exc:
                print(f"Base44 HTTP error: {exc.code}")
                self.send_json(502, {"error": "Could not retrieve transactions from Base44."})
            except Exception as exc:
                print(f"Report generation failed: {exc}")
                self.send_json(500, {"error": "Could not generate the Excel report."})
            return
        if path == "/":
            self.path = "/index.html"
            return super().do_GET()
        return super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/api/debts":
            self.send_json(404, {"error": "Not found"})
            return
        try:
            payload = self.read_json_body()
            name = str(payload.get("name", "")).strip()
            due_date_text = str(payload.get("due_date", "")).strip()
            reason = str(payload.get("reason", "")).strip()
            if not name or len(name) > 120:
                self.send_json(400, {"error": "Enter a name up to 120 characters."})
                return
            try:
                due_date = date.fromisoformat(due_date_text)
                amount = float(payload.get("amount"))
            except (ValueError, TypeError):
                self.send_json(400, {"error": "Enter a valid amount and due date."})
                return
            if not math.isfinite(amount) or amount <= 0:
                self.send_json(400, {"error": "Amount must be a positive finite number."})
                return
            debt = {
                "id": str(uuid4()), "name": name, "amount": round(amount, 2),
                "due_date": due_date.isoformat(), "reason": reason[:250],
                "created_at": datetime.now().isoformat(timespec="seconds"),
            }
            debts = load_debts()
            debts.append(debt)
            save_debts(debts)
            self.send_json(201, {"message": "Debt saved successfully.", "debt": prepare_debt(debt)})
        except (json.JSONDecodeError, UnicodeDecodeError, ValueError) as exc:
            self.send_json(400, {"error": str(exc)})
        except OSError as exc:
            print(f"Could not save debt: {exc}")
            self.send_json(500, {"error": "Could not save the debt."})
        except Exception as exc:
            print(f"Debt request failed: {exc}")
            self.send_json(500, {"error": "Could not process the debt."})

    def log_message(self, format, *args):
        print(f"{self.client_address[0]} - {format % args}")


if __name__ == "__main__":
    if not os.environ.get("BASE44_TOKEN"):
        raise SystemExit("BASE44_TOKEN is missing.")
    server = ThreadingHTTPServer((HOST, PORT), AppHandler)
    print(f"Server running on port {PORT}")
    try:
        server.serve_forever()
    finally:
        server.server_close()
